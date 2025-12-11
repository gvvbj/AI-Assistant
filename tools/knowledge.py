import os
import hashlib
import chromadb
from chromadb.api.types import Documents, EmbeddingFunction, Embeddings
from utils.logger import logger
from utils.error_handling import safe_execute
from tools.registry import tool_registry
import ollama
import re

# === 可选依赖导入 ===
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    from flashrank import Ranker, RerankRequest
    HAS_FLASHRANK = True
except ImportError:
    HAS_FLASHRANK = False
    logger.warning("FlashRank not installed. Rerank disabled.")

class OllamaEmbeddingFunction(EmbeddingFunction):
    def __init__(self, model_name="nomic-embed-text", base_url="http://127.0.0.1:11434"):
        self.model_name = model_name
        self.client = ollama.Client(host=base_url)

    def __call__(self, input: Documents) -> Embeddings:
        embeddings = []
        for text in input:
            try:
                resp = self.client.embeddings(model=self.model_name, prompt=text)
                embeddings.append(resp["embedding"])
            except Exception as e:
                logger.error(f"Embedding Error: {e}")
                embeddings.append([0.0]*768)
        return embeddings

class KnowledgeBase:
    _client = None
    _collection = None
    _current_embed_model = None
    _ranker = None
    _current_ranker_model = None

    def __init__(self):
        self.db_path = "chroma_db"
        os.makedirs(self.db_path, exist_ok=True)
        try:
            self._client = chromadb.PersistentClient(path=self.db_path)
        except Exception as e:
            logger.error(f"Chroma Init Fail: {e}")

    def _get_collection(self, embed_model_name):
        if not self._client: return None
        if self._collection is None or self._current_embed_model != embed_model_name:
            self._current_embed_model = embed_model_name
            try:
                safe_name = f"kb_{embed_model_name.replace(':', '_').replace('.', '_')}"
                self._collection = self._client.get_or_create_collection(
                    name=safe_name,
                    embedding_function=OllamaEmbeddingFunction(model_name=embed_model_name)
                )
            except Exception as e:
                logger.error(f"Collection Error: {e}")
                return None
        return self._collection

    def _get_ranker(self, model_name):
        if not HAS_FLASHRANK: return None
        real_name = model_name.split(" ")[0]
        if self._ranker is None or self._current_ranker_model != real_name:
            try:
                self._ranker = Ranker(model_name=real_name, cache_dir="models")
                self._current_ranker_model = real_name
            except Exception as e:
                logger.error(f"Reranker Init Fail: {e}")
                return None
        return self._ranker

    def _calculate_hash(self, file_path):
        h = hashlib.md5()
        with open(file_path, 'rb') as f: h.update(f.read())
        return h.hexdigest()

    def _extract_text(self, file_path):
        """提取文本：支持 txt, md, xlsx, pdf, docx"""
        ext = os.path.splitext(file_path)[1].lower()
        
        # 1. Excel
        if ext in ['.xlsx', '.xls']:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text = []
            for sheet in wb.worksheets:
                text.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    cleaned_row = [str(c) for c in row if c is not None]
                    if cleaned_row:
                        row_txt = " | ".join(cleaned_row)
                        text.append(row_txt)
            return "\n".join(text)
        
        # 2. PDF
        elif ext == '.pdf':
            if not pypdf: return "Error: 缺少 pypdf 库，无法解析 PDF"
            try:
                reader = pypdf.PdfReader(file_path)
                text = []
                for page in reader.pages:
                    text.append(page.extract_text() or "")
                return "\n".join(text)
            except Exception as e:
                return f"PDF解析失败: {e}"

        # 3. Word (Docx)
        elif ext == '.docx':
            if not docx: return "Error: 缺少 python-docx 库，无法解析 Docx"
            try:
                doc = docx.Document(file_path)
                return "\n".join([p.text for p in doc.paragraphs])
            except Exception as e:
                return f"Docx解析失败: {e}"
        
        # 4. 图片 (不支持)
        elif ext in ['.png', '.jpg', '.jpeg', '.bmp']:
            return None # 返回 None 表示不支持，上层会处理
            
        # 5. 纯文本
        elif ext in ['.txt', '.md', '.py', '.json', '.csv', '.html']:
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f: return f.read()
            except: return None
            
        return None

    # === 安全的迭代切分算法 (修复内存溢出/崩溃) ===
    def _safe_split_text(self, text, chunk_size=600, overlap=100):
        if not text: return []
        
        chunks = []
        total_len = len(text)
        start = 0
        
        while start < total_len:
            # 确定硬截止点
            end = min(start + chunk_size, total_len)
            
            # 如果还没到文件末尾，尝试优化切分点（找换行符）
            if end < total_len:
                # 在窗口后半部分寻找最近的换行符
                # 搜索范围：[end - chunk_size//2, end]
                lookback_limit = max(start, end - chunk_size // 2)
                
                # 优先找双换行（段落）
                last_newline = text.rfind('\n\n', lookback_limit, end)
                if last_newline != -1:
                    end = last_newline + 2 # 保留换行符
                else:
                    # 其次找单换行
                    last_newline = text.rfind('\n', lookback_limit, end)
                    if last_newline != -1:
                        end = last_newline + 1
                    else:
                        # 再其次找句号
                        last_period = text.rfind('。', lookback_limit, end)
                        if last_period != -1:
                            end = last_period + 1
                        
                        # 实在找不到分隔符，就硬切，不回退，防止死循环
            
            # 提取切片
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            
            # 计算下一个 start
            # 正常步进是 chunk长度 - overlap
            # 但要防止死循环（步进为0），强制至少前进 1
            step = max(1, (end - start) - overlap)
            
            # 如果是硬切且到了末尾，直接退出
            if end == total_len:
                break
                
            start += step
            
        return chunks

    @safe_execute("文档索引失败")
    def add_document(self, file_path, embed_model_name="nomic-embed-text"):
        coll = self._get_collection(embed_model_name)
        if not coll: return "DB连接失败"
        
        fname = os.path.basename(file_path)
        fhash = self._calculate_hash(file_path)
        
        # 检查是否已存在
        existing = coll.get(where={"file_hash": fhash})
        if existing['ids']:
            return f"文件 {fname} 已存在"
            
        content = self._extract_text(file_path)
        
        if content is None:
            # 特殊处理图片等不支持格式
            ext = os.path.splitext(file_path)[1].lower()
            if ext in ['.png', '.jpg', '.jpeg']:
                return "❌ 图片文件不支持文本索引。请使用'代码解释器'或 Vision 模型进行分析。"
            return f"❌ 格式 {ext} 不支持文本解析"
            
        if not content.strip(): 
            return "文件内容为空"
        
        # 使用新的安全切分策略
        chunks = self._safe_split_text(content, chunk_size=600, overlap=100)
        
        if not chunks: return "未能生成有效切片"

        ids = [f"{fhash}_{i}" for i in range(len(chunks))]
        metas = [{"source": fname, "file_hash": fhash} for _ in chunks]
        
        # 批量添加，防止单次请求过大
        batch_size = 100
        for i in range(0, len(chunks), batch_size):
            coll.add(
                documents=chunks[i:i+batch_size], 
                ids=ids[i:i+batch_size], 
                metadatas=metas[i:i+batch_size]
            )
            
        return f"索引成功，共生成 {len(chunks)} 个切片"

    @tool_registry.register(
        name="kb_search",
        description="Search the external Knowledge Base. Use this tool WHENEVER the user asks for information, facts, documents, or details that might be stored in the database.",
        parameters={
            "type": "object",
            "properties": {
                "query": {"type": "string", "description": "Keywords to search for"}
            },
            "required": ["query"]
        }
    )
    def search(self, query, embed_model_name="nomic-embed-text", rerank_model_name=None):
        coll = self._get_collection(embed_model_name)
        if not coll: return "DB Error"
        
        top_k = 15 if rerank_model_name else 5
        
        try:
            res = coll.query(query_texts=[query], n_results=top_k)
            docs = res['documents'][0]
            metas = res['metadatas'][0]
            
            if not docs: return "未找到相关内容"
            
            final_res = []
            
            if rerank_model_name and HAS_FLASHRANK:
                ranker = self._get_ranker(rerank_model_name)
                if ranker:
                    passages = [{"id": str(i), "text": d, "meta": m} for i, (d, m) in enumerate(zip(docs, metas))]
                    rerank_req = RerankRequest(query=query, passages=passages)
                    ranked_res = ranker.rerank(rerank_req)
                    
                    for item in ranked_res[:5]:
                        src = item['meta'].get('source', 'unknown')
                        final_res.append(f"📄 [Source: {src}]\n{item['text']}")
                    return "\n\n".join(final_res)

            for i in range(min(5, len(docs))):
                src = metas[i].get('source', 'unknown')
                final_res.append(f"📄 [Source: {src}]\n{docs[i]}")
            return "\n\n".join(final_res)
            
        except Exception as e:
            return f"检索异常: {e}"

    def get_files(self, embed_model_name="nomic-embed-text"):
        coll = self._get_collection(embed_model_name)
        if not coll: return []
        try:
            data = coll.get(include=['metadatas'])
            return list(set([m['source'] for m in data['metadatas'] if m]))
        except: return []

    def delete_file(self, fname, embed_model_name="nomic-embed-text"):
        coll = self._get_collection(embed_model_name)
        if coll: coll.delete(where={"source": fname})

knowledge_tool = KnowledgeBase()